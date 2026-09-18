"""Optional fixture regeneration: Python + openpyxl==3.1.5 (not a build dependency)."""
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904
root = Path(__file__).resolve().parents[1]
workbook = Workbook()
workbook.epoch = CALENDAR_MAC_1904
sheet = workbook.active
sheet.title = 'Data & notes'
sheet['A1'] = 'Independent producer'
sheet['A2'] = 10
sheet['A3'] = 20
sheet['B1'] = '=SUM(A2:A3)'
sheet['C1'] = True
sheet['D1'] = datetime(2024, 1, 1)
sheet['D1'].number_format = 'yyyy-mm-dd'
sheet['A1'].font = Font(bold=True, color='123456')
sheet['A1'].fill = PatternFill('solid', fgColor='EEEEAA')
summary = workbook.create_sheet('Résumé')
summary['A1'] = "='Data & notes'!B1*2"
summary['B1'] = '雪 & <text>'
workbook.save(root / 'tests/fixtures/openpyxl.xlsx')
# Verify the hand-authored OOXML document is also readable by an independent producer.
ours = load_workbook(root / 'tests/fixtures/basic.xlsx')
assert ours.sheetnames == ['Sales', 'Summary']
assert ours['Sales']['B6'].value == '=SUM(B2:B4)'
assert ours['Summary']['B1'].value == '=Sales!B6'
print('Independent OOXML fixture generated; basic.xlsx verified with openpyxl 3.1.5.')

from openpyxl.styles import Alignment
from openpyxl.worksheet.filters import FilterColumn, CustomFilters, CustomFilter
import shutil
layout = Workbook()
sheet = layout.active
sheet.title = 'Imported layout'
sheet.merge_cells('A1:D2')
sheet['A1'] = 'A workbook with its own layout\nWidths, merges, frozen headers and saved filters'
sheet['A1'].alignment = Alignment(wrap_text=True, vertical='center')
sheet['A1'].font = Font(bold=True, color='FFFFFF')
sheet['A1'].fill = PatternFill('solid', fgColor='254A39')
sheet.row_dimensions[1].height = 30
sheet.row_dimensions[2].height = 20
for column, width in [('A',24),('B',20),('C',32),('D',48),('T',20)]:
    sheet.column_dimensions[column].width = width
sheet.column_dimensions['C'].hidden = True
for column, value in zip('ABCD', ['Region','Revenue','Private notes','Description']):
    sheet[f'{column}4'] = value
    sheet[f'{column}4'].font = Font(bold=True)
for row, (region, revenue) in enumerate([('West',450),('East',300),('West',200),('East',100),('West',50)],5):
    sheet[f'A{row}'] = region
    sheet[f'B{row}'] = revenue
    sheet[f'C{row}'] = 'Hidden column'
    sheet[f'D{row}'] = 'A long description that stays inside its own column.'
sheet['T60'] = 'Last column'
sheet.row_dimensions[6].hidden = True
sheet.freeze_panes = 'B5'
sheet.auto_filter.ref = 'A4:D9'
sheet.auto_filter.add_filter_column(0, ['West'])
sheet.auto_filter.filterColumn.append(FilterColumn(colId=1,customFilters=CustomFilters(customFilter=[CustomFilter(operator='greaterThan',val='100')])))
sheet.auto_filter.add_sort_condition('B5:B9',descending=True)
layout.save(root / 'tests/fixtures/layout.xlsx')
shutil.copyfile(root / 'tests/fixtures/layout.xlsx',root / 'examples/react-demo/public/layout.xlsx')
print('Generated layout.xlsx with dimensions, merged cells, hidden dimensions, freeze panes and saved filters/sort.')
